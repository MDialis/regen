# app/services/ingestion/doc_processor.py
import base64
import csv
import io
import json
from dataclasses import dataclass, field
from pathlib import Path
from typing import List, Optional

from app.domain.ir_schema import (
    DataTypePrimitive,
    EntityDefinition,
    EntityField,
    SourceModality,
    SourceReference,
)
from app.services.ingestion.token_gatekeeper import token_gatekeeper


@dataclass
class ProcessedDocumentResult:
    """Output of deterministic Python preprocessing on non-AST files."""

    inferred_entities: List[EntityDefinition] = field(default_factory=list)
    condensed_summary_for_llm: str = ""
    raw_tokens_count: int = 0
    condensed_tokens_count: int = 0

    @property
    def tokens_saved(self) -> int:
        return max(0, self.raw_tokens_count - self.condensed_tokens_count)


class DocumentPreProcessor:
    """
    Pre-processes non-code files (.csv, .xlsx, .pdf, .json, .md, .txt) using Python
    so the LLM only reads compact structural digests instead of expensive raw dumps.
    """

    def process_file(
        self,
        file_path: str,
        text_content: Optional[str] = None,
        base64_content: Optional[str] = None,
    ) -> ProcessedDocumentResult:
        ext = Path(file_path).suffix.lower()
        stem = Path(file_path).stem.replace("-", "_").replace(" ", "_")
        pascal_name = "".join(w.capitalize() for w in stem.split("_") if w) or "Record"

        if ext == ".csv" and text_content:
            return self._process_csv(file_path, pascal_name, text_content)
        if ext == ".xlsx" and base64_content:
            return self._process_xlsx(file_path, pascal_name, base64_content)
        if ext == ".pdf" and base64_content:
            return self._process_pdf(file_path, base64_content)
        if ext == ".json" and text_content:
            return self._process_json(file_path, pascal_name, text_content)

        # Default: Markdown / Plain Text requirements document
        raw_text = text_content or ""
        raw_tokens = token_gatekeeper.count_tokens(raw_text)
        condensed = token_gatekeeper.truncate_to_budget(raw_text, 800)
        return ProcessedDocumentResult(
            condensed_summary_for_llm=f"[Document: {file_path}]\n{condensed}",
            raw_tokens_count=raw_tokens,
            condensed_tokens_count=token_gatekeeper.count_tokens(condensed),
        )

    def _process_csv(
        self, file_path: str, entity_name: str, content: str
    ) -> ProcessedDocumentResult:
        raw_tokens = token_gatekeeper.count_tokens(content)
        reader = csv.reader(io.StringIO(content))
        rows = []
        for idx, row in enumerate(reader):
            if row:
                rows.append(row)
            if idx >= 5:
                break

        if not rows:
            return ProcessedDocumentResult(raw_tokens_count=raw_tokens)

        headers = [h.strip() for h in rows[0] if h.strip()]
        sample_row = rows[1] if len(rows) > 1 else []
        fields = self._infer_fields_from_headers_and_sample(headers, sample_row)

        entity = EntityDefinition(
            name=entity_name,
            description=f"Inferred deterministically from CSV table `{file_path}`",
            fields=fields,
            sources=[
                SourceReference(
                    modality=SourceModality.PLAIN_TEXT_REQ,
                    file_path=file_path,
                    raw_symbol_name=entity_name,
                )
            ],
        )

        summary = (
            f"[CSV Schema Digest: {file_path}] -> Entity `{entity_name}` with columns: "
            + ", ".join(f"{f.name} ({f.data_type.value})" for f in fields)
        )
        return ProcessedDocumentResult(
            inferred_entities=[entity],
            condensed_summary_for_llm=summary,
            raw_tokens_count=raw_tokens,
            condensed_tokens_count=token_gatekeeper.count_tokens(summary),
        )

    def _process_xlsx(
        self, file_path: str, default_name: str, b64_data: str
    ) -> ProcessedDocumentResult:
        try:
            import openpyxl

            raw_bytes = base64.b64decode(b64_data)
            wb = openpyxl.load_workbook(
                io.BytesIO(raw_bytes), read_only=True, data_only=True
            )
            entities: List[EntityDefinition] = []
            summaries: List[str] = []

            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                rows = list(ws.iter_rows(values_only=True, max_row=4))
                if not rows or not rows[0]:
                    continue
                headers = [str(h).strip() for h in rows[0] if h is not None]
                sample = [
                    str(v) if v is not None else ""
                    for v in (rows[1] if len(rows) > 1 else [])
                ]
                ent_name = (
                    "".join(
                        w.capitalize() for w in sheet_name.replace("-", "_").split()
                    )
                    or default_name
                )
                fields = self._infer_fields_from_headers_and_sample(headers, sample)
                entities.append(
                    EntityDefinition(
                        name=ent_name,
                        description=f"Extracted from Excel sheet `{sheet_name}` in `{file_path}`",
                        fields=fields,
                        sources=[
                            SourceReference(
                                modality=SourceModality.PLAIN_TEXT_REQ,
                                file_path=file_path,
                                raw_symbol_name=sheet_name,
                            )
                        ],
                    )
                )
                summaries.append(
                    f"[Excel Sheet `{sheet_name}`] -> Entity `{ent_name}` columns: {', '.join(headers)}"
                )

            digest = "\n".join(summaries)
            return ProcessedDocumentResult(
                inferred_entities=entities,
                condensed_summary_for_llm=digest,
                raw_tokens_count=max(500, len(raw_bytes) // 8),
                condensed_tokens_count=token_gatekeeper.count_tokens(digest),
            )
        except Exception as exc:
            return ProcessedDocumentResult(
                condensed_summary_for_llm=f"[Excel file {file_path} could not be parsed: {exc}]"
            )

    def _process_pdf(self, file_path: str, b64_data: str) -> ProcessedDocumentResult:
        try:
            from pypdf import PdfReader

            raw_bytes = base64.b64decode(b64_data)
            reader = PdfReader(io.BytesIO(raw_bytes))
            extracted_pages = []
            for page in reader.pages[:10]:
                txt = page.extract_text() or ""
                if txt.strip():
                    extracted_pages.append(txt.strip())

            full_text = "\n".join(extracted_pages)
            raw_tokens = token_gatekeeper.count_tokens(full_text)
            condensed = token_gatekeeper.truncate_to_budget(full_text, 1000)
            summary = f"[PDF Requirements Digest: {file_path}]\n{condensed}"
            return ProcessedDocumentResult(
                condensed_summary_for_llm=summary,
                raw_tokens_count=raw_tokens,
                condensed_tokens_count=token_gatekeeper.count_tokens(summary),
            )
        except Exception as exc:
            return ProcessedDocumentResult(
                condensed_summary_for_llm=f"[PDF {file_path} parse note: {exc}]"
            )

    def _process_json(
        self, file_path: str, entity_name: str, content: str
    ) -> ProcessedDocumentResult:
        raw_tokens = token_gatekeeper.count_tokens(content)
        try:
            data = json.loads(content)
            sample = data[0] if isinstance(data, list) and data else data
            if isinstance(sample, dict):
                headers = list(sample.keys())
                values = [str(v) for v in sample.values()]
                fields = self._infer_fields_from_headers_and_sample(headers, values)
                entity = EntityDefinition(
                    name=entity_name,
                    description=f"Inferred from JSON structure in `{file_path}`",
                    fields=fields,
                    sources=[
                        SourceReference(
                            modality=SourceModality.PLAIN_TEXT_REQ,
                            file_path=file_path,
                        )
                    ],
                )
                summary = f"[JSON Schema Digest: {file_path}] -> Entity `{entity_name}` keys: {', '.join(headers)}"
                return ProcessedDocumentResult(
                    inferred_entities=[entity],
                    condensed_summary_for_llm=summary,
                    raw_tokens_count=raw_tokens,
                    condensed_tokens_count=token_gatekeeper.count_tokens(summary),
                )
        except Exception:
            pass
        condensed = token_gatekeeper.truncate_to_budget(content, 500)
        return ProcessedDocumentResult(
            condensed_summary_for_llm=f"[JSON File: {file_path}]\n{condensed}",
            raw_tokens_count=raw_tokens,
            condensed_tokens_count=token_gatekeeper.count_tokens(condensed),
        )

    @staticmethod
    def _infer_fields_from_headers_and_sample(
        headers: List[str], sample_values: List[str]
    ) -> List[EntityField]:
        fields: List[EntityField] = []
        for idx, raw_header in enumerate(headers):
            clean_name = (
                raw_header.strip()
                .lower()
                .replace(" ", "_")
                .replace("-", "_")
                .replace(".", "_")
            )
            if not clean_name:
                continue
            val = sample_values[idx].strip() if idx < len(sample_values) else ""
            dtype = DataTypePrimitive.STRING
            if val.isdigit():
                dtype = DataTypePrimitive.INTEGER
            else:
                try:
                    float(val)
                    dtype = DataTypePrimitive.FLOAT
                except ValueError:
                    if val.lower() in ("true", "false", "yes", "no"):
                        dtype = DataTypePrimitive.BOOLEAN

            fields.append(
                EntityField(
                    name=clean_name,
                    data_type=dtype,
                    is_required=True,
                    is_primary_key=clean_name in ("id", "uuid", "code", "pk"),
                )
            )
        return fields


doc_processor = DocumentPreProcessor()
